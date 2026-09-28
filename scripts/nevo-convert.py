#!/usr/bin/env python3
"""Zet het NEVO-bestand van het RIVM om naar src/data/nevo.json.

Gebruik: python3 scripts/nevo-convert.py NEVO2025_v9.0.csv [versie]
Werkt met de CSV (| of ; als scheidingsteken) of de Excel-versie (.xlsx, vereist openpyxl).
De waarden worden ongewijzigd overgenomen (voorwaarde van het RIVM).
"""
import csv, io, json, os, re, sys

OUT = os.path.join(os.path.dirname(__file__), '..', 'src', 'data', 'nevo.json')

# Kolom -> mogelijke kopteksten (NEVO gebruikt componentcodes als kolomnamen).
COLS = {
    'code': ['NEVO-code', 'NEVO-code/NEVO code', 'NEVO code'],
    'name': ['Voedingsmiddelnaam/Dutch food name', 'Voedingsmiddelnaam', 'Dutch food name'],
    'syn': ['Synoniem', 'Synoniem/Synonym', 'Synonym'],
    'qty': ['Hoeveelheid/Quantity', 'Hoeveelheid', 'Quantity'],
    'kcal': ['ENERCC'],
    'e': ['PROT'],
    'k': ['CHO'],
    'v': ['FAT'],
    'fiber': ['FIBT'],
    'sugar': ['SUGAR'],
    'satFat': ['FASAT'],
}


def rows_from(path):
    if path.lower().endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        # Zoek het blad met de NEVO-code-kolom.
        for sheet in wb.worksheets:
            first = next(sheet.iter_rows(max_row=1, values_only=True), ())
            if any(str(c or '').startswith('NEVO') for c in first):
                ws = sheet
                break
        return [[('' if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
    raw = open(path, 'rb').read()
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    head = text.split('\n', 1)[0]
    delim = max(['|', ';', '\t', ','], key=head.count)
    return list(csv.reader(io.StringIO(text), delimiter=delim))


def find(header, names):
    h = [str(x).strip() for x in header]
    for n in names:
        if n in h:
            return h.index(n)
    for n in names:  # kopteksten als "ENERCC (kcal)"
        for i, x in enumerate(h):
            if re.match(re.escape(n) + r'\b', x):
                return i
    return None


def num(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '.')
    if s in ('', '-', 'NA', 'n.a.'):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def r1(x):
    """Ongewijzigd; alleen 2.0 als 2 opslaan (scheelt ruimte)."""
    return None if x is None else (int(x) if float(x).is_integer() else x)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    rows = rows_from(path)
    hi = next(i for i, r in enumerate(rows[:20]) if find(r, COLS['code']) is not None and find(r, COLS['kcal']) is not None)
    header = rows[hi]
    idx = {k: find(header, v) for k, v in COLS.items()}
    missing = [k for k, v in idx.items() if v is None and k not in ('syn', 'fiber', 'sugar', 'satFat', 'qty')]
    if missing:
        sys.exit(f'Kolommen niet gevonden: {missing}\nKoppen: {header}')

    get = lambda r, k: r[idx[k]] if idx[k] is not None and idx[k] < len(r) else ''
    items, skipped = [], 0
    for r in rows[hi + 1:]:
        code = num(get(r, 'code'))
        name = str(get(r, 'name')).strip()
        kcal, e, k, v = (num(get(r, x)) for x in ('kcal', 'e', 'k', 'v'))
        if code is None or not name or None in (kcal, e, k, v):
            skipped += 1
            continue
        unit = 'ml' if 'ml' in str(get(r, 'qty')).lower() else 'g'
        syn = str(get(r, 'syn') or '').strip()
        items.append([int(code), name, syn, unit, r1(kcal), r1(e), r1(k), r1(v),
                      r1(num(get(r, 'fiber'))), r1(num(get(r, 'sugar'))), r1(num(get(r, 'satFat')))])

    m = re.search(r'(20\d\d)[^0-9]*v?(\d+(?:\.\d+)?)', os.path.basename(path))
    version = sys.argv[2] if len(sys.argv) > 2 else (f'{m.group(1)}/{m.group(2)}' if m else '')
    data = {
        'version': version,
        'source': f'NEVO-online versie {version}, RIVM, Bilthoven'.replace('  ', ' '),
        'items': sorted(items, key=lambda x: x[0]),
    }
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, separators=(',', ':'))
    print(f'{len(items)} producten opgeslagen ({skipped} overgeslagen), versie {version}, '
          f'{os.path.getsize(OUT) // 1024} kB -> {os.path.normpath(OUT)}')


if __name__ == '__main__':
    main()
